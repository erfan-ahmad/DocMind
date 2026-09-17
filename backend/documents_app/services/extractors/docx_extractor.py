import logging
from io import BytesIO

from docx import Document as DocxDocument

from .persian_utils import fix_persian_text
from ..exceptions import TextExtractionError

logger = logging.getLogger(__name__)


def extract_from_docx(file_bytes: bytes) -> str:
    """استخراج متن از فایل Word"""
    try:
        doc = DocxDocument(BytesIO(file_bytes))
    except Exception as exc:
        logger.error("python-docx خطا داد: %s", exc)
        raise TextExtractionError(f"خطا در باز کردن Word: {exc}")

    text_parts = []

    # ۱) پاراگراف‌ها
    for paragraph in doc.paragraphs:
        if paragraph.text.strip():
            text_parts.append(paragraph.text)

    # ۲) جدول‌ها
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            text_parts.append(' | '.join(cells))

    text = '\n'.join(text_parts).strip()

    if not text:
        raise TextExtractionError("هیچ متنی از Word استخراج نشد.")

    return fix_persian_text(text)